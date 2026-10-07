# -*- coding: utf-8 -*-
"""Critérium Fédéral — NATIONALE 1 : classeur de la FFTT -> criterium2627.json

Le tableau N1 est un classeur Google Sheets public, un onglet par catégorie et par
groupe (JGA, JGB, SMA, SMB…). Export :
  curl -L -o n1.xlsx "https://docs.google.com/spreadsheets/d/<id>/export?format=xlsx"

⚠️ Le classeur garde des onglets d'ANCIENNES SAISONS (CF, CG, JF, JG, SD, SM portent
2023/2024). On ne retient que ceux dont l'en-tête annonce la saison demandée, sinon on
ferait jouer des joueurs partis depuis longtemps.

Usage : python3 scripts/criterium_n1_xlsx.py <tour> <saison> <salle> <classeur.xlsx>
Exemple : … 1 2026/2027 "Joué-lès-Tours (37)" /tmp/n1.xlsx
"""
import json, os, re, sys, unicodedata, datetime
import openpyxl

SORTIE = 'data/criterium2627.json'
PRIVE = 'data/_criterium_lic.json'
CLUB_ACBB = '08920049'
MOIS = {'janvier':1,'fevrier':2,'mars':3,'avril':4,'mai':5,'juin':6,'juillet':7,
        'aout':8,'septembre':9,'octobre':10,'novembre':11,'decembre':12}

CAT = {'CF': ('-15', 'F', 'Cadettes'), 'CG': ('-15', 'M', 'Cadets'),
       'JF': ('-19', 'F', 'Juniors Filles'), 'JG': ('-19', 'M', 'Juniors Garçons'),
       'SD': ('ELITE', 'F', 'Seniors Dames'), 'SM': ('ELITE', 'M', 'Seniors Messieurs')}
ORDRE = ['-11', '-13', '-15', '-19', 'ELITE']
LIB = {('ELITE', 'F'): 'Élite Dames', ('ELITE', 'M'): 'Élite Messieurs'}

def coupe(bloc):
    """« MOREL-GONZALES Titouan » -> nom, prénom (le nom est la tête en capitales)."""
    mots = str(bloc or '').split()
    def caps(m):
        return unicodedata.normalize('NFD', m).encode('ascii', 'ignore').decode().replace('-', '').isupper()
    i = 0
    while i < len(mots) and caps(mots[i]) and len(mots[i]) > 1:
        i += 1
    if i == 0: return str(bloc or ''), ''
    return ' '.join(mots[:i]), ' '.join(mots[i:])

def points(cls, officiel):
    for v in (officiel, cls):
        m = re.findall(r'(\d{3,4})', str(v or ''))
        if m: return int(m[-1])
    return 0

def main():
    if len(sys.argv) < 5: sys.exit('usage : criterium_n1_xlsx.py <tour> <saison> <salle> <classeur.xlsx>')
    tour, saison, salle, chemin = sys.argv[1], sys.argv[2], sys.argv[3], os.path.expanduser(sys.argv[4])
    wb = openpyxl.load_workbook(chemin, data_only=True)
    groupes, ignores = [], []
    for ws in wb.worksheets:
        # Les licences arrivent en nombres décimaux (9454825.0) : str() puis « chiffres seulement » donnait
        # 94548250, un zéro de trop — les 4 joueurs du club en N1 ne pouvaient pas confirmer (07/10/2026).
        cell = lambda c: '' if c is None else (str(int(c)) if isinstance(c, float) and c.is_integer() else str(c)).strip()
        lignes = [[cell(c) for c in r] for r in ws.iter_rows(values_only=True)]
        tete = ' '.join(x for r in lignes[:7] for x in r if x)
        if saison not in tete:
            ignores.append(ws.title); continue
        base = re.sub(r'[^A-Z]', '', ws.title.upper())[:2]
        if base not in CAT: ignores.append(ws.title); continue
        age, genre, lib = CAT[base]
        suffixe = ws.title.upper()[2:]                    # A, B… = groupe de niveau
        date = re.search(r'(\d{1,2}),?\s*(\d{1,2})\s*et\s*(\d{1,2})\s+(\w+)\s+(\d{4})', tete)
        g = {'nom': 'N1 ' + lib + (' ' + suffixe if suffixe else ''), 'niveau': 'N', 'div': 'N1',
             'age': age, 'genre': genre,
             'cat': LIB.get((age, genre)) or (age + ' ans ' + ('Filles' if genre == 'F' else 'Garçons')),
             'salle': salle, 'date': '', 'pointage': '', 'debut': '', 'contact': '', 'joueurs': []}
        if date:
            mois = unicodedata.normalize('NFD', date.group(4)).encode('ascii', 'ignore').decode().lower()
            if mois in MOIS:
                # Le classeur annonce un week-end, pas un jour unique par joueur.
                g['date'] = datetime.date(int(date.group(5)), MOIS[mois], int(date.group(1))).strftime('%d/%m/%Y')
                g['date_fin'] = datetime.date(int(date.group(5)), MOIS[mois], int(date.group(3))).strftime('%d/%m/%Y')
        for r in lignes:
            if not r or not re.fullmatch(r'\d+(\.0)?', str(r[0]).strip()): continue
            pos = int(float(r[0])); lic = re.sub(r'\D', '', str(r[1] if len(r) > 1 else ''))
            nom, pre = coupe(r[2] if len(r) > 2 else '')
            club = str(r[6] if len(r) > 6 else '')
            numclub, _, nomclub = club.partition(' - ')
            if not lic or not nom: continue
            acbb = numclub.strip() == CLUB_ACBB
            j = {'pos': pos, 'ajoute': False, 'nom': nom, 'pre': pre, 'club': (nomclub or club).strip(),
                 'cat': str(r[5] if len(r) > 5 else '')[:12], 'cats': '',
                 'clt': points(r[3] if len(r) > 3 else '', r[7] if len(r) > 7 else ''), 'acbb': acbb}
            if acbb: j['lic'] = lic
            j['_lic'] = lic
            g['joueurs'].append(j)
        if g['joueurs']: groupes.append(g)
    if ignores: print('  onglets ignorés (autre saison ou hors format) : %s' % ', '.join(ignores))

    doc = json.load(open(SORTIE)); t = doc['tours'][str(tour)]
    # Le tour s'étale sur un week-end : la N1 commence le vendredi, le départemental
    # jeunes joue le samedi, les adultes le dimanche. Un libellé de plage vaut mieux
    # qu'une date unique, qui serait fausse pour la plupart.
    for ws in wb.worksheets:
        tete = ' '.join(str(c) for r in ws.iter_rows(min_row=1, max_row=7, values_only=True) for c in r if c)
        if saison not in tete: continue
        m = re.search(r'(\d{1,2}),?\s*(\d{1,2})\s*et\s*(\d{1,2})\s+(\w+)\s+(\d{4})', tete)
        if m:
            t['weekend'] = '%s, %s et %s %s %s' % m.groups(); break
    t['groupes'] = [x for x in t['groupes'] if not (x.get('niveau') == 'N' and x.get('div') == 'N1')] + groupes
    for g in t['groupes']:
        g.setdefault('niveau', 'D')
        g['rang'] = [ORDRE.index(g['age']) if g.get('age') in ORDRE else 9,
                     0 if g.get('genre') == 'F' else 1,
                     {'D': 0, 'R': 1, 'N': 2}.get(g['niveau'], 9)]
    for i, g in enumerate(t['groupes']): g['id'] = 't%s-g%02d' % (tour, i + 1)
    # licences hors du fichier publié, comme pour les autres échelons
    table = json.load(open(PRIVE)) if os.path.exists(PRIVE) else {}
    for g in t['groupes']:
        for j in g['joueurs']:
            lic = j.pop('_lic', None)
            if lic: table.setdefault(g['id'], {})[str(j.get('pos') or 0)] = lic
    json.dump(table, open(PRIVE, 'w'), ensure_ascii=False, indent=1, sort_keys=True)
    doc['maj'] = datetime.date.today().isoformat()
    json.dump(doc, open(SORTIE, 'w'), ensure_ascii=False, indent=1, sort_keys=False)
    acbb = sum(1 for g in groupes for j in g['joueurs'] if j['acbb'])
    print('%s : N1 · %d groupes · %d joueurs dont %d ACBB' % (SORTIE, len(groupes), sum(len(g['joueurs']) for g in groupes), acbb))

if __name__ == '__main__':
    main()
